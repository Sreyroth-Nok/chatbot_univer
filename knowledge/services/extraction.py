"""Read plain text out of uploaded PDF, Word and text files."""
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def extract_text(file_obj, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(file_obj)
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    if ext == ".docx":
        import docx

        document = docx.Document(file_obj)
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n\n".join(parts)
    if ext in (".txt", ".md"):
        data = file_obj.read()
        return data.decode("utf-8", errors="replace") if isinstance(data, bytes) else data
    raise ValueError(
        f"Unsupported file type '{ext}'. Upload one of: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
    )
